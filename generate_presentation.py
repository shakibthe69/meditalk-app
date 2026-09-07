from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt


ROOT = Path(__file__).parent
OUT = ROOT / "docs" / "Meditalk_Product_Presentation.pptx"
ASSETS = ROOT / "docs" / "presentation_assets"
ASSETS.mkdir(parents=True, exist_ok=True)

W, H = 768, 1365
BG = "#F6F9FA"
WHITE = "#FFFFFF"
INK = "#122033"
MUTED = "#66768A"
TEAL = "#129A90"
TEAL_DARK = "#116B69"
TEAL_LIGHT = "#DDF8F3"
MINT = "#EFFCF9"
BLUE = "#3A72D8"
BLUE_LIGHT = "#EAF1FF"
GREEN = "#20B86B"
RED = "#EF5B5B"
ORANGE = "#F4A524"
LINE = "#E2E9EE"


def font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/Aptos.ttf" if not bold else "C:/Windows/Fonts/Aptos-Bold.ttf",
        "C:/Windows/Fonts/Arial.ttf" if not bold else "C:/Windows/Fonts/Arialbd.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return ImageFont.truetype(path, size)
    return ImageFont.load_default()


def hex_rgb(value):
    value = value.lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def rounded(draw, box, radius, fill, outline=None, width=1):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def text(draw, xy, value, size, fill=INK, bold=False, anchor=None):
    draw.text(xy, value, font=font(size, bold), fill=fill, anchor=anchor)


def wrap(draw, xy, value, max_width, size, fill=INK, bold=False, gap=8):
    words = value.split()
    lines, current = [], ""
    f = font(size, bold)
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=f)[2] <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    y = xy[1]
    for line in lines:
        draw.text((xy[0], y), line, font=f, fill=fill)
        y += size + gap
    return y


def icon(draw, x, y, symbol, color=TEAL, size=28):
    text(draw, (x, y), symbol, size, color, True, anchor="mm")


def phone_base(draw, x=70, y=35, width=628, height=1295):
    rounded(draw, (x, y, x + width, y + height), 48, WHITE, "#D5E0E6", 4)
    rounded(draw, (x + 214, y + 14, x + width - 214, y + 38), 16, INK)
    return x + 28, y + 64, width - 56, height - 95


def header(draw, x, y, title, subtitle, add=True):
    text(draw, (x, y), title, 34, INK, True)
    text(draw, (x, y + 44), subtitle, 19, MUTED)
    draw.line((x, y + 86, x + 572, y + 86), fill=LINE, width=2)
    if add:
        rounded(draw, (x + 510, y - 7, x + 570, y + 53), 30, TEAL)
        text(draw, (x + 540, y + 23), "+", 34, WHITE, True, anchor="mm")


def search(draw, x, y, label):
    rounded(draw, (x, y, x + 572, y + 70), 34, WHITE, LINE, 2)
    text(draw, (x + 24, y + 35), "?", 24, "#91A0B2", False, anchor="mm")
    text(draw, (x + 52, y + 35), label, 19, "#91A0B2", anchor="lm")


def pill(draw, x, y, label, active=False, width=None):
    width = width or max(80, len(label) * 10 + 34)
    rounded(draw, (x, y, x + width, y + 38), 19, TEAL if active else WHITE, TEAL if active else LINE, 2)
    text(draw, (x + width / 2, y + 19), label, 16, WHITE if active else INK, True, anchor="mm")


def medicine_card(draw, x, y, name, generic, dose, schedule, instruction, active=True):
    rounded(draw, (x, y, x + 572, y + 260), 28, WHITE, LINE, 2)
    rounded(draw, (x + 26, y + 25, x + 84, y + 83), 29, MINT, TEAL_LIGHT, 2)
    icon(draw, x + 55, y + 54, "+", TEAL_DARK, 25)
    text(draw, (x + 102, y + 32), name, 26, INK, True)
    text(draw, (x + 102, y + 66), generic, 17, MUTED)
    pill(draw, x + 456, y + 32, "ACTIVE" if active else "DONE", True, 92)
    rounded(draw, (x + 26, y + 105, x + 546, y + 150), 14, BG)
    text(draw, (x + 42, y + 119), "DOSAGE", 13, "#9AA7B7", True)
    text(draw, (x + 42, y + 136), dose, 19, INK, True)
    text(draw, (x + 308, y + 119), "FREQUENCY", 13, "#9AA7B7", True)
    text(draw, (x + 308, y + 136), "DAILY", 19, INK, True)
    text(draw, (x + 30, y + 177), f"o  Schedules: {schedule}", 17, MUTED, True)
    rounded(draw, (x + 26, y + 207, x + 546, y + 242), 12, MINT)
    text(draw, (x + 42, y + 225), f"*  {instruction}", 15, TEAL_DARK, anchor="lm")


def make_dashboard():
    im = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(im)
    x, y, sw, sh = phone_base(d)
    text(d, (x, y + 6), "Good day, SHAKIB", 29, INK, True)
    text(d, (x, y + 42), "[]  Monday, Sep 7, 2026", 16, MUTED, True)
    rounded(d, (x + 16, y + 88, x + sw - 16, y + 172), 24, MINT, TEAL_LIGHT, 2)
    icon(d, x + 50, y + 130, "*", TEAL, 25)
    text(d, (x + 82, y + 110), "UPCOMING MEDICATION", 14, TEAL_DARK, True)
    text(d, (x + 82, y + 140), "Napa 500mg at 02:00 PM", 19, INK, True)
    rounded(d, (x + 16, y + 198, x + sw - 16, y + 395), 28, WHITE, LINE, 2)
    text(d, (x + 32, y + 226), "Medication Adherence", 23, INK, True)
    pill(d, x + 360, y + 216, "90%+ GOAL", False, 155)
    d.ellipse((x + 45, y + 260, x + 190, y + 405), outline=TEAL, width=14)
    text(d, (x + 118, y + 322), "100%", 27, INK, True, anchor="mm")
    text(d, (x + 118, y + 350), "ADHERENCE", 12, MUTED, True, anchor="mm")
    for idx, (label, value, color) in enumerate([("Taken", "100%", GREEN), ("Missed", "0%", RED), ("Skipped", "0%", ORANGE)]):
        yy = y + 270 + idx * 42
        d.ellipse((x + 235, yy, x + 249, yy + 14), fill=color)
        text(d, (x + 266, yy - 1), label, 17, MUTED)
        text(d, (x + 510, yy - 1), value, 17, INK, True, anchor="ra")
    for idx, (sym, label, color) in enumerate([("+", "Add Med", TEAL), ("[]", "Scan Rx", BLUE), ("[]", "Add Report", GREEN)]):
        xx = x + 16 + idx * 187
        rounded(d, (xx, y + 425, xx + 170, y + 535), 22, WHITE, LINE, 2)
        rounded(d, (xx + 56, y + 440, xx + 114, y + 498), 29, TEAL_LIGHT if idx == 0 else (BLUE_LIGHT if idx == 1 else "#EAF9EF"))
        icon(d, xx + 85, y + 469, sym, color, 25)
        text(d, (xx + 85, y + 515), label, 16, INK, True, anchor="mm")
    text(d, (x + 16, y + 580), "Today's Schedule", 24, INK, True)
    text(d, (x + sw - 16, y + 588), "View All", 17, TEAL, True, anchor="ra")
    rounded(d, (x + 16, y + 620, x + sw - 16, y + 830), 28, WHITE, LINE, 2)
    text(d, (x + 32, y + 648), "o  02:00 PM", 18, INK, True)
    pill(d, x + 440, y + 638, "PENDING", False, 110)
    text(d, (x + 32, y + 690), "Napa 500mg", 25, INK, True)
    text(d, (x + 32, y + 725), "Dosage: 1 Tablet", 17, MUTED)
    rounded(d, (x + 32, y + 760, x + 250, y + 796), 12, MINT)
    text(d, (x + 48, y + 778), "*  After lunch / meal", 15, TEAL_DARK, anchor="lm")
    rounded(d, (x + 32, y + 805, x + 208, y + 850), 22, BG, LINE, 2)
    text(d, (x + 120, y + 827), "x  Skip", 16, INK, True, anchor="mm")
    rounded(d, (x + 224, y + 805, x + 540, y + 850), 22, TEAL)
    text(d, (x + 382, y + 827), "OK  Take Medicine", 16, WHITE, True, anchor="mm")
    return im


def make_medicines():
    im = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(im)
    x, y, sw, sh = phone_base(d)
    header(d, x, y + 5, "My Medicines", "2 total recorded medications")
    search(d, x, y + 110, "Search medicine or generic name...")
    pill(d, x, y + 130 + 68, "All Medications", True, 174)
    pill(d, x + 184, y + 198, "Active", False, 88)
    pill(d, x + 282, y + 198, "Completed", False, 120)
    medicine_card(d, x, y + 260, "Omeprazole", "Omeprazole 20mg Capsule", "20mg", "08:00 AM (1 Capsule)", "Take 30 minutes before breakfast")
    medicine_card(d, x, y + 535, "Napa", "Paracetamol 500mg", "500mg", "02:00 PM, 10:00 PM", "Take after meals for fever and pain")
    return im


def make_history():
    im = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(im)
    x, y, sw, sh = phone_base(d)
    header(d, x, y + 5, "Medical History", "Chronological health journey")
    search(d, x, y + 110, "Search records, doctors, test names...")
    pill(d, x, y + 198, "All Records", True, 158)
    pill(d, x + 168, y + 198, "Prescriptions", False, 142)
    pill(d, x + 320, y + 198, "Lab Reports", False, 122)
    text(d, (x + 18, y + 258), "[]  TIMELINE RECORDS", 16, TEAL_DARK, True)
    d.line((x + 32, y + 322, x + 32, y + 1040), fill="#A5EBDD", width=5)
    cards = [
        ("OK", "Sept 12, 2026", "Treatment Completed", "Napa 500mg (10 days course)", "COMPLETED"),
        ("~", "Sept 6, 2026", "Complete Blood Count (CBC)", "National Diagnostic Lab", "VERIFIED REPORT"),
        ("[]", "Sept 5, 2026", "Prescription Digitalized", "3 medicines prescribed", "PRESCRIPTION"),
    ]
    for idx, (sym, date, title, sub, badge) in enumerate(cards):
        yy = y + 305 + idx * 240
        d.ellipse((x + 2, yy, x + 62, yy + 60), fill=WHITE, outline=TEAL, width=3)
        icon(d, x + 32, yy + 30, sym, TEAL, 23)
        rounded(d, (x + 82, yy, x + 572, yy + 210), 26, WHITE, LINE, 2)
        text(d, (x + 110, yy + 28), date, 16, TEAL_DARK, True)
        text(d, (x + 110, yy + 60), title, 22, INK, True)
        text(d, (x + 110, yy + 94), sub, 17, MUTED)
        pill(d, x + 378, yy + 30, badge, False, 160)
        wrap(d, (x + 110, yy + 132), "Verified and organized for your next clinical conversation.", 400, 15, MUTED)
    return im


def make_ocr():
    im = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(im)
    x, y, sw, sh = phone_base(d)
    header(d, x, y + 5, "Add Prescription", "Capture, extract, verify")
    wrap(d, (x, y + 115), "Turn a prescription photo into an editable medication plan.", 560, 19, MUTED)
    for idx, (title, sub, symbol, color) in enumerate([("Take Photo", "Capture using camera", "[]", TEAL), ("Upload Image", "Select from gallery", "[]", BLUE)]):
        xx = x + idx * 294
        rounded(d, (xx, y + 205, xx + 270, y + 330), 25, WHITE, LINE, 2)
        rounded(d, (xx + 90, y + 222, xx + 180, y + 290), 34, TEAL_LIGHT if idx == 0 else BLUE_LIGHT)
        icon(d, xx + 135, y + 256, symbol, color, 30)
        text(d, (xx + 135, y + 307), title, 19, INK, True, anchor="mm")
        text(d, (xx + 135, y + 331), sub, 14, MUTED, anchor="mm")
    text(d, (x + 286, y + 384), "OR QUICK DEMO", 14, "#9AA7B7", True, anchor="mm")
    d.line((x, y + 384, x + 200, y + 384), fill=LINE, width=2)
    d.line((x + 372, y + 384, x + 572, y + 384), fill=LINE, width=2)
    rounded(d, (x, y + 430, x + 572, y + 505), 25, MINT, "#9DEDE0", 3)
    icon(d, x + 45, y + 468, "*", TEAL, 28)
    text(d, (x + 82, y + 451), "Use Sample Prescription", 20, TEAL_DARK, True)
    text(d, (x + 82, y + 480), "Instant OCR demo with patient verification", 15, TEAL_DARK)
    rounded(d, (x, y + 555, x + 572, y + 820), 26, WHITE, LINE, 2)
    text(d, (x + 24, y + 582), "AI Medicine Review", 22, INK, True)
    pill(d, x + 385, y + 580, "REVIEW", False, 120)
    for idx, line in enumerate(["Omeprazole 20mg   1x daily", "Napa 500mg         2x daily", "Rosuvastatin 10mg  1x nightly"]):
        rounded(d, (x + 24, y + 638 + idx * 48, x + 548, y + 676 + idx * 48), 12, BG)
        text(d, (x + 42, y + 646 + idx * 48), line, 16, INK, True)
    rounded(d, (x, y + 850, x + 572, y + 900), 25, TEAL)
    text(d, (x + 286, y + 875), "Confirm and Create Reminders", 17, WHITE, True, anchor="mm")
    return im


def make_report():
    im = Image.new("RGB", (W, H), BG)
    d = ImageDraw.Draw(im)
    x, y, sw, sh = phone_base(d)
    header(d, x, y + 5, "Upload Medical Report", "Keep every result in one place", add=False)
    fields = ["Report / Test Title *", "Diagnostic Lab / Hospital", "Ordering Physician", "Clinical Findings & Notes"]
    values = ["Complete Blood Count (CBC)", "National Diagnostic Lab", "Dr. Rahman, MD", "WBC, RBC, platelet and hemoglobin within baseline"]
    for idx, (label, value) in enumerate(zip(fields, values)):
        yy = y + 125 + idx * 125
        text(d, (x, yy), label, 17, INK, True)
        rounded(d, (x, yy + 28, x + 572, yy + 82), 27, WHITE, LINE, 2)
        text(d, (x + 24, yy + 55), value, 16, INK, anchor="lm")
    text(d, (x, y + 640), "Test Category", 17, INK, True)
    pill(d, x, y + 670, "Blood Test", True, 126)
    pill(d, x + 138, y + 670, "X-Ray", False, 90)
    pill(d, x + 238, y + 670, "MRI Scan", False, 108)
    rounded(d, (x, y + 760, x + 572, y + 875), 25, WHITE, "#9DEDE0", 3)
    icon(d, x + 286, y + 807, "^", TEAL, 32)
    text(d, (x + 286, y + 846), "Choose report image / PDF", 17, INK, True, anchor="mm")
    rounded(d, (x, y + 910, x + 572, y + 962), 26, TEAL)
    text(d, (x + 286, y + 936), "Save Medical Report", 18, WHITE, True, anchor="mm")
    return im


def save_screen(name, maker):
    path = ASSETS / f"{name}.png"
    maker().save(path)
    return path


def add_bg(slide, color="F6F9FA"):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = hex_rgb(color)


def add_text(slide, x, y, w, h, value, size=20, color=INK, bold=False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = value
    run.font.name = "Aptos Display"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = hex_rgb(color)
    return box


def add_shape(slide, shape_type, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = hex_rgb(fill)
    shape.line.color.rgb = hex_rgb(line or fill)
    return shape


def add_label(slide, x, y, label, number, color=TEAL):
    add_shape(slide, MSO_SHAPE.OVAL, x, y, 0.34, 0.34, color)
    add_text(slide, x, y + 0.01, 0.34, 0.25, str(number), 11, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, x + 0.46, y - 0.02, 2.6, 0.4, label, 14, INK, True)


def add_phone(slide, image, x, y, w=3.0):
    slide.shapes.add_picture(str(image), Inches(x), Inches(y), width=Inches(w))


def build_deck():
    screens = {
        "dashboard": save_screen("dashboard", make_dashboard),
        "medicines": save_screen("medicines", make_medicines),
        "history": save_screen("history", make_history),
        "ocr": save_screen("ocr", make_ocr),
        "report": save_screen("report", make_report),
    }
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Cover
    slide = prs.slides.add_slide(blank); add_bg(slide, "0E3E43")
    add_shape(slide, MSO_SHAPE.OVAL, 9.4, -1.2, 5.2, 5.2, "146D6D")
    add_shape(slide, MSO_SHAPE.OVAL, 10.2, 4.8, 3.8, 3.8, "0B5155")
    add_text(slide, 0.8, 0.65, 7, 0.35, "MEDITALK HEALTHCARE PLATFORM", 15, TEAL_LIGHT, True)
    add_text(slide, 0.8, 1.55, 7.2, 1.7, "Health records\nthat move with you.", 31, WHITE, True)
    add_text(slide, 0.8, 3.55, 6.7, 0.9, "A simple, secure companion for prescriptions, reminders, reports, and doctor-ready history.", 17, "C4E8E3")
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, 5.45, 3.1, 0.55, TEAL)
    add_text(slide, 0.8, 5.58, 3.1, 0.25, "PRODUCT PRESENTATION  |  2026", 12, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, 0.8, 6.55, 5.2, 0.5, "MD Shakib Ahmed  |  Tarek Rahman  |  Mayesha Islam", 12, "C4E8E3")
    add_text(slide, 11.0, 6.55, 1.4, 0.5, "01", 12, TEAL_LIGHT, True, PP_ALIGN.RIGHT)

    # 2. Product promise
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_text(slide, 0.75, 0.55, 8, 0.45, "One calm place for a complex health journey", 25, INK, True)
    add_text(slide, 0.75, 1.14, 8, 0.45, "Meditalk turns scattered medical moments into one clear daily routine.", 15, MUTED)
    cards = [("01", "Remember", "Upcoming doses and simple Take / Skip actions", TEAL), ("02", "Record", "Prescriptions, reports, and visits stay searchable", BLUE), ("03", "Share", "Export a clean summary for the next doctor visit", GREEN)]
    for idx, (num, title, desc, color) in enumerate(cards):
        xx = 0.75 + idx * 4.12
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, xx, 2.1, 3.65, 2.35, WHITE, LINE)
        add_shape(slide, MSO_SHAPE.OVAL, xx + 0.28, 2.42, 0.58, 0.58, color)
        add_text(slide, xx + 0.28, 2.58, 0.58, 0.2, num, 12, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, xx + 0.28, 3.2, 3, 0.35, title, 21, INK, True)
        add_text(slide, xx + 0.28, 3.72, 3.0, 0.55, desc, 14, MUTED)
    add_text(slide, 0.75, 5.15, 5.5, 0.35, "DESIGNED FOR THE MOMENT THAT MATTERS", 12, TEAL_DARK, True)
    add_text(slide, 0.75, 5.65, 11.3, 0.75, "When the patient is busy, the prescription is handwritten, and the doctor needs the full story.", 25, TEAL_DARK, True)

    # 3. Dashboard
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_text(slide, 0.7, 0.45, 7.2, 0.45, "The dashboard answers: what do I do next?", 25, INK, True)
    add_text(slide, 0.7, 1.03, 6.5, 0.35, "A glanceable home screen for today's care routine.", 15, MUTED)
    add_phone(slide, screens["dashboard"], 0.7, 1.45, 3.45)
    add_label(slide, 4.65, 1.85, "Upcoming dose", 1)
    add_text(slide, 5.12, 2.22, 6.7, 0.45, "The next medication is visible before the user has to search.", 15, MUTED)
    add_label(slide, 4.65, 3.0, "Adherence ring", 2)
    add_text(slide, 5.12, 3.37, 6.7, 0.45, "Taken, missed, and skipped doses become an immediate signal.", 15, MUTED)
    add_label(slide, 4.65, 4.15, "Quick actions", 3)
    add_text(slide, 5.12, 4.52, 6.7, 0.45, "Add Med, Scan Rx, and Add Report are one tap away.", 15, MUTED)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 4.65, 5.45, 7.2, 0.8, MINT, TEAL_LIGHT)
    add_text(slide, 4.95, 5.67, 6.6, 0.35, "CLICK FLOW  →  Take Medicine  →  adherence updates", 15, TEAL_DARK, True)

    # 4. Medicines
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_text(slide, 0.7, 0.45, 7, 0.45, "Medication management without the clutter", 25, INK, True)
    add_text(slide, 0.7, 1.03, 7, 0.35, "Search, filter, schedule, and safely remove active treatments.", 15, MUTED)
    add_phone(slide, screens["medicines"], 0.7, 1.35, 3.45)
    add_label(slide, 4.65, 1.95, "Search + filters", 1)
    add_text(slide, 5.12, 2.32, 6.7, 0.45, "Find a medicine by brand or generic name, then narrow by status.", 15, MUTED)
    add_label(slide, 4.65, 3.1, "Schedule clarity", 2)
    add_text(slide, 5.12, 3.47, 6.7, 0.45, "Dose, frequency, time, and food instruction sit together.", 15, MUTED)
    add_label(slide, 4.65, 4.25, "Safe controls", 3)
    add_text(slide, 5.12, 4.62, 6.7, 0.45, "Status toggles and delete confirmation reduce accidental changes.", 15, MUTED)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 4.65, 5.55, 7.2, 0.8, "FFF7E8", "F9DDA2")
    add_text(slide, 4.95, 5.77, 6.6, 0.35, "PATIENT-FIRST DETAIL  |  before / after meal guidance", 15, "8A5B06", True)

    # 5. OCR flow
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_text(slide, 0.7, 0.45, 7.8, 0.45, "From prescription photo to verified plan", 25, INK, True)
    add_text(slide, 0.7, 1.03, 8.5, 0.35, "OCR accelerates entry; the patient remains in control before anything is saved.", 15, MUTED)
    add_phone(slide, screens["ocr"], 0.7, 1.35, 3.45)
    steps = [("Capture", "Camera or gallery image"), ("Extract", "OCR text and medicine details"), ("Verify", "Edit before creating reminders")]
    for idx, (title, desc) in enumerate(steps):
        yy = 2.0 + idx * 1.28
        add_shape(slide, MSO_SHAPE.OVAL, 4.8, yy, 0.62, 0.62, TEAL if idx < 2 else GREEN)
        add_text(slide, 4.8, yy + 0.2, 0.62, 0.2, str(idx + 1), 13, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, 5.7, yy + 0.02, 2.7, 0.3, title, 19, INK, True)
        add_text(slide, 5.7, yy + 0.38, 5.8, 0.35, desc, 15, MUTED)
        if idx < 2:
            add_shape(slide, MSO_SHAPE.RECTANGLE, 5.08, yy + 0.62, 0.06, 0.66, "A7E8DE")
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 4.8, 6.05, 7.1, 0.75, MINT, TEAL_LIGHT)
    add_text(slide, 5.1, 6.26, 6.5, 0.3, "SAFETY GATE  |  no schedule is committed without confirmation", 14, TEAL_DARK, True)

    # 6. History and reports
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_text(slide, 0.7, 0.45, 8.5, 0.45, "A health story that is easy to scan", 25, INK, True)
    add_text(slide, 0.7, 1.03, 8.5, 0.35, "Timeline records connect treatments, tests, prescriptions, and visits.", 15, MUTED)
    add_phone(slide, screens["history"], 0.7, 1.35, 3.45)
    add_phone(slide, screens["report"], 4.45, 1.35, 3.45)
    add_label(slide, 8.15, 2.0, "Filter the story", 1)
    add_text(slide, 8.62, 2.37, 3.9, 0.45, "Search by doctor, test, or record type.", 14, MUTED)
    add_label(slide, 8.15, 3.25, "Add evidence", 2)
    add_text(slide, 8.62, 3.62, 3.9, 0.45, "Upload report image or PDF with notes.", 14, MUTED)
    add_label(slide, 8.15, 4.5, "Keep it clinical", 3)
    add_text(slide, 8.62, 4.87, 3.9, 0.65, "Doctor, facility, category, and findings stay attached to the record.", 14, MUTED)

    # 7. Export / architecture
    slide = prs.slides.add_slide(blank); add_bg(slide, "103C42")
    add_text(slide, 0.75, 0.55, 8.5, 0.45, "Ready for the next consultation", 25, WHITE, True)
    add_text(slide, 0.75, 1.13, 8.5, 0.38, "Meditalk turns everyday tracking into a doctor-ready summary.", 15, "C4E8E3")
    layers = [("MOBILE", "React Native + Expo\nTypeScript + Zustand", TEAL), ("SERVICES", "Spring Boot\nJWT + OCR + OpenPDF", BLUE), ("DATA", "MySQL + JPA\nUser-isolated records", GREEN)]
    for idx, (title, body, color) in enumerate(layers):
        xx = 0.75 + idx * 4.15
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, xx, 2.1, 3.55, 1.75, "164E54", "2D6B70")
        add_shape(slide, MSO_SHAPE.RECTANGLE, xx, 2.1, 3.55, 0.14, color)
        add_text(slide, xx + 0.25, 2.48, 3, 0.3, title, 13, color, True)
        add_text(slide, xx + 0.25, 2.95, 3, 0.65, body, 17, WHITE, True)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.75, 4.55, 11.85, 1.05, "E4F9F5", "7EDBCE")
    add_text(slide, 1.1, 4.82, 11.1, 0.35, "PDF EXPORT  →  Profile  +  Active Medicines  +  Prescriptions  +  Lab Results", 17, TEAL_DARK, True, PP_ALIGN.CENTER)
    add_text(slide, 0.75, 6.35, 11.5, 0.4, "Safety boundary: Meditalk organizes records and reminders; it does not replace medical advice.", 14, "C4E8E3", False, PP_ALIGN.CENTER)

    # 8. Click journey
    slide = prs.slides.add_slide(blank); add_bg(slide)
    add_text(slide, 0.7, 0.45, 7.5, 0.45, "A simple click journey", 25, INK, True)
    add_text(slide, 0.7, 1.03, 8.5, 0.35, "Every primary action has a visible next step and a clear outcome.", 15, MUTED)
    journey = [("Home", "See next dose", TEAL), ("Take", "Log adherence", GREEN), ("History", "Find the record", BLUE), ("Export", "Share the summary", ORANGE)]
    for idx, (title, desc, color) in enumerate(journey):
        xx = 0.9 + idx * 3.05
        add_shape(slide, MSO_SHAPE.OVAL, xx, 2.2, 1.05, 1.05, color)
        add_text(slide, xx, 2.57, 1.05, 0.25, str(idx + 1), 22, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, xx - 0.2, 3.55, 1.45, 0.3, title, 18, INK, True, PP_ALIGN.CENTER)
        add_text(slide, xx - 0.65, 3.97, 2.35, 0.45, desc, 14, MUTED, False, PP_ALIGN.CENTER)
        if idx < 3:
            add_text(slide, xx + 1.28, 2.55, 0.6, 0.35, "→", 28, "#9AA7B7", True, PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.3, 5.2, 10.7, 0.95, BG, LINE)
    add_text(slide, 1.55, 5.5, 10.2, 0.3, "The product stays useful because it keeps the user's next decision obvious.", 18, TEAL_DARK, True, PP_ALIGN.CENTER)

    # 9. Team
    slide = prs.slides.add_slide(blank); add_bg(slide, "0E3E43")
    add_text(slide, 0.8, 0.72, 7, 0.45, "Built by the Meditalk team", 27, WHITE, True)
    add_text(slide, 0.8, 1.32, 8.2, 0.4, "A focused product for safer, more organized everyday care.", 16, "C4E8E3")
    team = [("MD SHAKIB AHMED", "ID: 0112230862"), ("TAREK RAHMAN", "ID: 0000000000000"), ("MAYESHA ISLAM", "ID: 000000000")]
    for idx, (name, ident) in enumerate(team):
        yy = 2.35 + idx * 1.1
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.8, yy, 8.4, 0.82, "164E54", "2D6B70")
        add_shape(slide, MSO_SHAPE.OVAL, 1.08, yy + 0.18, 0.45, 0.45, TEAL)
        add_text(slide, 1.78, yy + 0.2, 5.4, 0.27, name, 15, WHITE, True)
        add_text(slide, 6.3, yy + 0.2, 2.55, 0.27, ident, 13, "C4E8E3", False, PP_ALIGN.RIGHT)
    add_text(slide, 0.8, 6.35, 11.4, 0.4, "Meditalk  |  Secure records. Clear reminders. Better conversations.", 17, TEAL_LIGHT, True)

    # 10. Close
    slide = prs.slides.add_slide(blank); add_bg(slide, "F1FBF9")
    add_text(slide, 0.85, 1.3, 11.7, 0.75, "A clearer health journey\nstarts with one small action.", 31, TEAL_DARK, True, PP_ALIGN.CENTER)
    add_text(slide, 1.2, 3.15, 10.9, 0.6, "Remember the dose. Record the moment. Bring the full story.", 19, MUTED, False, PP_ALIGN.CENTER)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 4.55, 4.45, 4.25, 0.72, TEAL)
    add_text(slide, 4.55, 4.68, 4.25, 0.28, "MEDITALK", 19, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, 0.85, 6.65, 11.7, 0.3, "Thank you", 14, TEAL_DARK, True, PP_ALIGN.CENTER)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_deck()